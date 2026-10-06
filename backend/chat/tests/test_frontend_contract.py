"""Walks through every call of the SPA's src/api/http.js, like the browser does.

Session + CSRF are enforced (the token is re-read from the cookie before each unsafe
request, since login rotates it) and RealocAI is the fake client, so the report
blocks are the real examples. The asserted keys are the ones the SPA's mock client
returns (src/api/mock/mockClient.js), which is the executable spec of the contract.
"""

import io
import re

import pytest
from django.contrib.auth import get_user_model
from openpyxl import load_workbook
from rest_framework.test import APIClient

pytestmark = pytest.mark.django_db

EMAIL = "equipe@clinica.test"
PASSWORD = "senha-forte-123"
USER_KEYS = {"id", "email", "full_name", "is_staff"}
PAGE_KEYS = {"count", "next", "previous", "results"}
CONVERSATION_KEYS = {
    "id",
    "title",
    "status",
    "summary",
    "message_count",
    "created_at",
    "updated_at",
}
MESSAGE_KEYS = {"id", "seq", "role", "content", "created_at"}
# Same regex as requestBlob() in http.js.
FILENAME = re.compile(r"filename\*?=(?:UTF-8'')?\"?([^\";]+)\"?", re.IGNORECASE)


class Browser:
    """Mimics http.js: JSON bodies, cookies, X-CSRFToken on unsafe methods."""

    def __init__(self):
        self.client = APIClient(enforce_csrf_checks=True)

    def request(self, method, path, body=None, query=None):
        headers = {"HTTP_ACCEPT": "application/json"}
        if method in ("POST", "PUT", "PATCH", "DELETE"):
            token = self.client.cookies.get("csrftoken")
            if token:
                headers["HTTP_X_CSRFTOKEN"] = token.value
        call = getattr(self.client, method.lower())
        if method == "GET":
            return call(path, query or {}, **headers)
        return (
            call(path, body, format="json", **headers)
            if body is not None
            else call(path, **headers)
        )


def assert_error(response, status):
    assert response.status_code == status
    assert set(response.json()) == {"detail"}
    assert isinstance(response.json()["detail"], str)


@pytest.fixture
def fake_realocai(settings):
    settings.REALOCAI_USE_FAKE = True


def test_every_http_js_call(fake_realocai):
    get_user_model().objects.create_user(email=EMAIL, password=PASSWORD, full_name="Equipe")
    b = Browser()

    # me() before login → 401 (the SPA redirects to /login).
    assert_error(b.request("GET", "/api/auth/me/"), 401)

    # csrf()
    response = b.request("GET", "/api/auth/csrf/")
    assert response.status_code == 200
    assert b.client.cookies["csrftoken"].value

    # login(email, password)
    assert_error(b.request("POST", "/api/auth/login/", {"email": EMAIL, "password": "x"}), 401)
    response = b.request("POST", "/api/auth/login/", {"email": EMAIL, "password": PASSWORD})
    assert response.status_code == 200
    assert set(response.json()) == USER_KEYS

    # me()
    response = b.request("GET", "/api/auth/me/")
    assert response.status_code == 200
    assert set(response.json()) == USER_KEYS

    # createConversation(title) — the SPA sends {} when there is no title.
    response = b.request("POST", "/api/conversations/", {})
    assert response.status_code == 201
    conversation = response.json()
    assert set(conversation) == CONVERSATION_KEYS
    assert (conversation["status"], conversation["summary"], conversation["message_count"]) == (
        "active",
        "",
        0,
    )
    cid = conversation["id"]
    response = b.request("POST", "/api/conversations/", {"title": "Com título"})
    assert response.json()["title"] == "Com título"

    # listConversations({page, status})
    for status in (None, "active", "archived"):
        query = {"page": 1, **({"status": status} if status else {})}
        response = b.request("GET", "/api/conversations/", query=query)
        assert response.status_code == 200
        assert set(response.json()) == PAGE_KEYS
        for item in response.json()["results"]:
            assert set(item) == CONVERSATION_KEYS

    # getConversation(id)
    response = b.request("GET", f"/api/conversations/{cid}/")
    assert response.status_code == 200
    assert set(response.json()) == CONVERSATION_KEYS

    # sendMessage(id, content) — empty first, then a report request.
    assert_error(b.request("POST", f"/api/conversations/{cid}/messages/", {"content": " "}), 400)
    response = b.request(
        "POST", f"/api/conversations/{cid}/messages/", {"content": "Relatório completo"}
    )
    assert 200 <= response.status_code < 300
    body = response.json()
    assert set(body) == {"user_message", "assistant_message"}
    assert set(body["user_message"]) == MESSAGE_KEYS
    assert set(body["assistant_message"]) == MESSAGE_KEYS | {"blocos"}
    blocos = body["assistant_message"]["blocos"]
    assert [b_["tipo"] for b_ in blocos] == [
        "ocupacao_profissional",
        "pacientes_por_profissional",
        "ocupacao_agregada",
    ]

    # The sidebar gets the automatic title and the preview.
    conversation = b.request("GET", f"/api/conversations/{cid}/").json()
    assert conversation["title"] == "Relatório completo"
    assert conversation["summary"] == body["assistant_message"]["content"]
    assert conversation["message_count"] == 2

    # listMessages(id, {page}) — the history reopens with the cards.
    response = b.request("GET", f"/api/conversations/{cid}/messages/", query={"page": 1})
    assert response.status_code == 200
    page = response.json()
    assert set(page) == PAGE_KEYS
    assert [m["role"] for m in page["results"]] == ["user", "assistant"]
    assert page["results"][1]["blocos"] == blocos
    assert_error(b.request("GET", f"/api/conversations/{cid}/messages/", query={"page": 9}), 404)

    # exportReports({blocos, formato})
    response = b.request("POST", "/api/reports/export/", {"formato": "excel", "blocos": blocos})
    assert response.status_code == 200
    filename = FILENAME.search(response["Content-Disposition"]).group(1)
    assert re.fullmatch(r"realocai-relatorios-\d{4}-\d{2}-\d{2}\.xlsx", filename)
    assert load_workbook(io.BytesIO(response.content)).sheetnames[0] == "Resumo"
    assert_error(
        b.request("POST", "/api/reports/export/", {"formato": "pdf", "blocos": blocos}), 400
    )

    # updateConversation(id, patch): rename, archive; archived conversations are read-only.
    response = b.request("PATCH", f"/api/conversations/{cid}/", {"title": "Renomeada"})
    assert response.status_code == 200
    assert set(response.json()) == CONVERSATION_KEYS
    response = b.request("PATCH", f"/api/conversations/{cid}/", {"status": "archived"})
    assert response.json()["status"] == "archived"
    assert_error(b.request("PATCH", f"/api/conversations/{cid}/", {"status": "x"}), 400)
    assert_error(b.request("POST", f"/api/conversations/{cid}/messages/", {"content": "Oi"}), 400)

    # deleteConversation(id)
    response = b.request("DELETE", f"/api/conversations/{cid}/")
    assert response.status_code == 204
    assert_error(b.request("GET", f"/api/conversations/{cid}/"), 404)

    # logout()
    assert b.request("POST", "/api/auth/logout/").status_code == 204
    assert_error(b.request("GET", "/api/auth/me/"), 401)
    assert_error(b.request("GET", "/api/conversations/"), 401)
