"""Excel (.xlsx) export of report blocks, following "Especificação do Excel" in the
SPA's docs/relatorios.md (reference implementation: gerarExcel in localExport.js).

- "Resumo" sheet: one row per block (Relatório, Período, Meta, Gerado em, Avisos);
- one sheet per item of `tabelas`, columns in `colunas` order;
- bold header on #E6F3EF, first row frozen, widths by format/content;
- values are written as they came: never recalculated.

Formula injection: every string is stored as a text cell (data_type "s"), so a value
such as "=cmd|..." is never evaluated by Excel.
"""

import datetime
import io
import re

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

SUMMARY_SHEET = "Resumo"
SUMMARY_COLUMNS = [
    ("Relatório", "texto"),
    ("Período", "texto"),
    ("Meta", "percentual"),
    ("Gerado em", "data"),
    ("Avisos", "texto"),
]
PARTIAL_DATA_NOTICE = "Dados parciais"

NUMBER_FORMATS = {
    "percentual": "0.0%",
    "inteiro": "0",
    "decimal": "0.0",
    "data": "dd/mm/yyyy",
}
NUMERIC_FORMATS = ("percentual", "inteiro", "decimal")
FIXED_WIDTHS = {"percentual": 10, "inteiro": 10, "decimal": 10, "data": 12}
TEXT_MIN_WIDTH = 22
TEXT_MAX_WIDTH = 40

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill(fill_type="solid", start_color="E6F3EF", end_color="E6F3EF")

SHEET_NAME_MAX_LENGTH = 31
_FORBIDDEN_SHEET_CHARS = re.compile(r"[\\/?*\[\]:]")
EXCEL_CELL_MAX_LENGTH = 32767

# datetime.date.weekday(): Monday is 0.
_WEEKDAYS = ["Segunda", "Terça", "Quarta", "Quinta", "Sexta", "Sábado", "Domingo"]


def export_filename(blocos, today):
    """realocai-<tipo>-<AAAA-MM-DD>.xlsx, or realocai-relatorios-... for several blocks.

    `tipo` was validated against ^[a-z0-9_]{1,50}$, so the name is plain ASCII.
    """
    tipo = blocos[0]["tipo"] if len(blocos) == 1 else "relatorios"
    return f"realocai-{tipo}-{today.isoformat()}.xlsx"


def format_period(periodo):
    """Same text as periodo() in format.js: "28/09 a 03/10/2026" or "Segunda, 28/09/2026"."""
    inicio = datetime.date.fromisoformat(periodo["inicio"])
    fim = datetime.date.fromisoformat(periodo["fim"])
    if inicio == fim:
        return f"{_WEEKDAYS[inicio.weekday()]}, {inicio:%d/%m/%Y}"
    return f"{inicio:%d/%m} a {fim:%d/%m/%Y}"


def build_workbook(blocos, generated_on):
    """Return the .xlsx bytes for already validated blocks."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = SUMMARY_SHEET
    used_names = {SUMMARY_SHEET.lower()}

    rows = [
        [
            bloco["titulo"],
            format_period(bloco["periodo"]),
            bloco["meta"],
            generated_on,
            " | ".join([*bloco["avisos"], *([PARTIAL_DATA_NOTICE] if bloco["parcial"] else [])]),
        ]
        for bloco in blocos
    ]
    _fill_sheet(summary, SUMMARY_COLUMNS, rows)

    several = len(blocos) > 1
    for index, bloco in enumerate(blocos, start=1):
        for tabela in bloco["tabelas"]:
            suffix = f" ({index})" if several else ""
            sheet = workbook.create_sheet(sheet_name(tabela["nome"], suffix, used_names))
            columns = [(c["rotulo"], c["formato"]) for c in tabela["colunas"]]
            keys = [c["chave"] for c in tabela["colunas"]]
            rows = [[linha.get(key) for key in keys] for linha in tabela["linhas"]]
            _fill_sheet(sheet, columns, rows)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def sheet_name(nome, suffix, used_names):
    """Sheet name: no \\ / ? * [ ] :, at most 31 characters, unique (case-insensitive)."""
    base = _FORBIDDEN_SHEET_CHARS.sub(" ", nome)
    base = ILLEGAL_CHARACTERS_RE.sub("", base).strip().strip("'").strip() or "Tabela"
    name = _fit(base, suffix)
    candidate, counter = name, 2
    while candidate.lower() in used_names:
        candidate = _fit(name, f" {counter}")
        counter += 1
    used_names.add(candidate.lower())
    return candidate


def _fit(base, suffix):
    return base[: SHEET_NAME_MAX_LENGTH - len(suffix)].rstrip() + suffix


def _fill_sheet(sheet, columns, rows):
    widths = [len(label) for label, _ in columns]
    for col, (label, _) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=col)
        _set_text(cell, label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    for row_index, values in enumerate(rows, start=2):
        for col, ((_, formato), value) in enumerate(zip(columns, values, strict=True), start=1):
            shown = _write_value(sheet.cell(row=row_index, column=col), value, formato)
            widths[col - 1] = max(widths[col - 1], len(shown))

    for col, ((_, formato), content_width) in enumerate(zip(columns, widths, strict=True), 1):
        width = FIXED_WIDTHS.get(formato)
        if width is None:
            width = min(max(content_width + 2, TEXT_MIN_WIDTH), TEXT_MAX_WIDTH)
        sheet.column_dimensions[get_column_letter(col)].width = width

    sheet.freeze_panes = "A2"


def _write_value(cell, value, formato):
    """Write one value; return its text, used to size text columns."""
    if value is None or value == "":
        return ""  # empty cell, never "—"
    if formato in NUMERIC_FORMATS and _is_number(value):
        cell.value = value
        cell.number_format = NUMBER_FORMATS[formato]
        return str(value)
    if formato == "data":
        date = _as_date(value)
        if date is not None:
            cell.value = date
            cell.number_format = NUMBER_FORMATS["data"]
            return "00/00/0000"
    text = _as_text(value)
    _set_text(cell, text)
    return text


def _set_text(cell, text):
    text = ILLEGAL_CHARACTERS_RE.sub("", text)[:EXCEL_CELL_MAX_LENGTH]
    cell.value = text
    # openpyxl turns strings starting with "=" into formulas; force a plain string so
    # "=cmd|...", "+...", "-...", "@..." are never evaluated.
    cell.data_type = "s"


def _is_number(value):
    return isinstance(value, int | float) and not isinstance(value, bool)


def _as_date(value):
    if isinstance(value, datetime.date):
        return value
    if isinstance(value, str):
        try:
            return datetime.date.fromisoformat(value[:10]) if len(value) >= 10 else None
        except ValueError:
            return None
    return None


def _as_text(value):
    if isinstance(value, bool):
        return "Sim" if value else "Não"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)
