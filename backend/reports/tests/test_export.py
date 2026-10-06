import copy
import datetime
import io
import json
from pathlib import Path

import pytest
from django.contrib.auth import get_user_model
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from chat.models import Conversation
from chat.services.conversations import record_exchange
from reports.excel import XLSX_CONTENT_TYPE

pytestmark = pytest.mark.django_db

FIXTURES = Path(__file__).resolve().parents[2] / "chat" / "tests" / "fixtures"
URL = "/api/reports/export/"


def load(name):
    return json.loads((FIXTURES / f"{name}.json").read_text(encoding="utf-8"))


@pytest.fixture
def user(db):
    return get_user_model().objects.create_user(
        email="profissional@exemplo.com", password="senha-forte-123"
    )


@pytest.fixture
def api(user):
    client = APIClient()
    client.force_login(user)
    return client


def export(client, blocos, formato="excel"):
    return client.post(URL, {"formato": formato, "blocos": blocos}, format="json")


def open_xlsx(response):
    assert response.status_code == 200, response.content[:300]
    return load_workbook(io.BytesIO(response.content))


def minimal_block(**overrides):
    bloco = {
        "versao": 1,
        "tipo": "ocupacao_profissional",
        "titulo": "Ocupação fictícia",
        "periodo": {"inicio": "2026-10-05", "fim": "2026-10-10"},
        "meta": 0.8,
        "parcial": False,
        "avisos": [],
        "resumo": [],
        "dados": {"qualquer": ["coisa", {"aninhada": True}]},
        "tabelas": [
            {
                "nome": "Por dia",
                "colunas": [
                    {"chave": "data", "rotulo": "Data", "formato": "data"},
                    {"chave": "nome", "rotulo": "Nome", "formato": "texto"},
                    {"chave": "pct", "rotulo": "Ocupação", "formato": "percentual"},
                    {"chave": "qtd", "rotulo": "Slots", "formato": "inteiro"},
                    {"chave": "media", "rotulo": "Média", "formato": "decimal"},
                ],
                "linhas": [
                    {"data": "2026-10-05", "nome": "Sala 1", "pct": 0.75, "qtd": 3, "media": 1.5}
                ],
            }
        ],
    }
    bloco.update(overrides)
    return bloco


def rows(sheet):
    return [[c.value for c in row] for row in sheet.iter_rows()]


# --- Authentication, CSRF and format --------------------------------------------------


def test_requires_session():
    response = export(APIClient(), [minimal_block()])

    assert response.status_code == 401
    assert response.json() == {"detail": "Não autenticado."}


def test_requires_csrf(user):
    client = APIClient(enforce_csrf_checks=True)
    client.force_login(user)

    response = export(client, [minimal_block()])

    assert response.status_code == 403
    assert set(response.json()) == {"detail"}


def test_works_with_csrf_token(user):
    client = APIClient(enforce_csrf_checks=True)
    client.get(reverse("auth:csrf"))
    client.force_login(user)
    token = client.cookies["csrftoken"].value

    response = client.post(
        URL,
        {"formato": "excel", "blocos": [minimal_block()]},
        format="json",
        HTTP_X_CSRFTOKEN=token,
    )

    assert response.status_code == 200


@pytest.mark.parametrize("formato", ["pdf", "csv", "", None, "EXCEL"])
def test_unsupported_format(api, formato):
    response = export(api, [minimal_block()], formato=formato)

    assert response.status_code == 400
    assert response.json() == {"detail": "Formato não suportado."}


def test_missing_format(api):
    response = api.post(URL, {"blocos": [minimal_block()]}, format="json")

    assert response.json() == {"detail": "Formato não suportado."}


# --- Strict validation -----------------------------------------------------------------


@pytest.mark.parametrize("count", [0, 11])
def test_rejects_block_count_outside_1_to_10(api, count):
    response = export(api, [minimal_block()] * count)

    assert response.status_code == 400
    assert response.json() == {"detail": "Envie de 1 a 10 blocos."}


def test_accepts_ten_blocks(api):
    assert export(api, [minimal_block()] * 10).status_code == 200


def test_rejects_malformed_json(api):
    response = api.post(URL, "{formato: excel", content_type="application/json")

    assert response.status_code == 400
    assert response.json() == {"detail": "Requisição malformada."}


@pytest.mark.parametrize("body", [[], "texto", 3])
def test_rejects_body_that_is_not_an_object(api, body):
    response = api.post(URL, json.dumps(body), content_type="application/json")

    assert response.status_code == 400
    assert set(response.json()) == {"detail"}


@pytest.mark.parametrize(
    "change",
    [
        {"tipo": "Ocupação"},
        {"tipo": "a" * 51},
        {"tipo": "../etc"},
        {"tipo": 3},
        {"titulo": 42},
        {"meta": "0.8"},
        {"meta": True},
        {"parcial": "false"},
        {"parcial": 0},
        {"periodo": {"inicio": "05/10/2026", "fim": "2026-10-10"}},
        {"periodo": {"inicio": "2026-02-30", "fim": "2026-10-10"}},
        {"avisos": "um aviso"},
        {"avisos": [1]},
        {"resumo": [{"rotulo": "x", "valor": {"obj": 1}, "formato": "texto"}]},
        {"tabelas": "nada"},
        {"tabelas": [{"nome": "T", "colunas": [], "linhas": []}]},
        {
            "tabelas": [
                {
                    "nome": "T",
                    "colunas": [{"chave": "a", "rotulo": "A", "formato": "moeda"}],
                    "linhas": [],
                }
            ]
        },  # noqa: E501
        {
            "tabelas": [
                {
                    "nome": "T",
                    "colunas": [{"chave": "a", "rotulo": "A", "formato": "texto"}],
                    "linhas": [{"a": [1]}],
                }
            ]
        },  # noqa: E501
        {
            "tabelas": [
                {
                    "nome": "T",
                    "colunas": [{"chave": "a", "rotulo": "A", "formato": "texto"}],
                    "linhas": ["x"],
                }
            ]
        },  # noqa: E501
    ],
)
def test_rejects_invalid_fields(api, change):
    response = export(api, [minimal_block(**change)])

    assert response.status_code == 400
    assert set(response.json()) == {"detail"}


@pytest.mark.parametrize("field", ["tipo", "titulo", "periodo", "meta", "parcial", "tabelas"])
def test_rejects_missing_fields(api, field):
    bloco = minimal_block()
    del bloco[field]

    assert export(api, [bloco]).status_code == 400


def test_meta_may_be_null_and_percentages_above_one_pass(api):
    bloco = minimal_block(meta=None)
    bloco["tabelas"][0]["linhas"][0]["pct"] = 1.25

    workbook = open_xlsx(export(api, [bloco]))

    assert workbook["Resumo"]["C2"].value is None
    assert workbook["Por dia"]["C2"].value == 1.25


def test_dados_and_unknown_keys_are_ignored(api):
    bloco = minimal_block(dados="não validado", extra={"x": 1}, versao=99)

    assert export(api, [bloco]).status_code == 200


def test_body_above_the_limit_is_413(api, settings):
    settings.REPORTS_EXPORT_MAX_BYTES = 1000
    bloco = minimal_block(avisos=["x" * 900])

    response = export(api, [bloco])

    assert response.status_code == 413
    assert set(response.json()) == {"detail"}


# --- Spreadsheet content (real examples from RealocAI) ----------------------------------


def test_single_real_block(api):
    bloco = load("ocupacao_profissional")
    today = timezone.localdate()

    response = export(api, [bloco])
    workbook = open_xlsx(response)

    assert response["Content-Type"] == XLSX_CONTENT_TYPE
    assert response["Content-Disposition"] == (
        f'attachment; filename="realocai-ocupacao_profissional-{today:%Y-%m-%d}.xlsx"'
    )
    assert workbook.sheetnames == ["Resumo", *[t["nome"] for t in bloco["tabelas"]]]

    summary = workbook["Resumo"]
    assert rows(summary)[0] == ["Relatório", "Período", "Meta", "Gerado em", "Avisos"]
    assert summary["A2"].value == bloco["titulo"]
    assert summary["B2"].value == "28/09 a 03/10/2026"
    assert (summary["C2"].value, summary["C2"].number_format) == (0.8, "0.0%")
    assert summary["D2"].value == datetime.datetime.combine(today, datetime.time())
    assert summary["D2"].number_format == "dd/mm/yyyy"
    assert summary["E2"].value == " | ".join(bloco["avisos"])


def test_table_cells_follow_the_column_formats(api):
    bloco = load("ocupacao_profissional")
    tabela = bloco["tabelas"][0]

    sheet = open_xlsx(export(api, [bloco]))[tabela["nome"]]

    assert [c.value for c in sheet[1]] == [c["rotulo"] for c in tabela["colunas"]]
    assert sheet.max_row == len(tabela["linhas"]) + 1
    expected_formats = {
        "percentual": "0.0%",
        "inteiro": "0",
        "decimal": "0.0",
        "data": "dd/mm/yyyy",
    }
    for col, coluna in enumerate(tabela["colunas"], start=1):
        for row, linha in enumerate(tabela["linhas"], start=2):
            cell = sheet.cell(row=row, column=col)
            value = linha.get(coluna["chave"])
            if coluna["formato"] == "data":
                assert cell.value == datetime.datetime.fromisoformat(value)
            elif coluna["formato"] in ("percentual", "inteiro", "decimal"):
                assert cell.value == value
            if coluna["formato"] in expected_formats and value is not None:
                assert cell.number_format == expected_formats[coluna["formato"]]


def test_header_style_frozen_row_and_widths(api):
    workbook = open_xlsx(export(api, [minimal_block()]))

    for sheet in workbook.worksheets:
        assert sheet.freeze_panes == "A2"
        for cell in sheet[1]:
            assert cell.font.bold
            assert cell.fill.fill_type == "solid"
            assert cell.fill.start_color.rgb.endswith("E6F3EF")

    widths = {k: v.width for k, v in workbook["Por dia"].column_dimensions.items()}
    # data, texto, percentual, inteiro, decimal
    assert widths == {"A": 12, "B": 22, "C": 10, "D": 10, "E": 10}
    summary_widths = {k: v.width for k, v in workbook["Resumo"].column_dimensions.items()}
    assert (summary_widths["C"], summary_widths["D"]) == (10, 12)


def test_text_width_grows_with_content_up_to_40(api):
    bloco = minimal_block()
    bloco["tabelas"][0]["linhas"].append({"nome": "x" * 30})
    bloco["tabelas"][0]["linhas"].append({"nome": "y" * 80})
    medium = minimal_block()
    medium["tabelas"][0]["linhas"][0]["nome"] = "z" * 30

    assert open_xlsx(export(api, [bloco]))["Por dia"].column_dimensions["B"].width == 40
    assert open_xlsx(export(api, [medium]))["Por dia"].column_dimensions["B"].width == 32


def test_empty_values_are_empty_cells(api):
    bloco = minimal_block()
    bloco["tabelas"][0]["linhas"] = [{"data": None, "nome": "", "pct": None}]

    sheet = open_xlsx(export(api, [bloco]))["Por dia"]

    assert [c.value for c in sheet[2]] == [None, None, None, None, None]


def test_booleans_in_text_columns_become_sim_nao(api):
    bloco = load("ocupacao_agregada")

    sheet = open_xlsx(export(api, [bloco]))["Por especialidade"]

    assert sheet["E2"].value == "Sim"


def test_single_day_period_and_null_meta(api):
    bloco = minimal_block(periodo={"inicio": "2026-09-28", "fim": "2026-09-28"}, meta=None)

    summary = open_xlsx(export(api, [bloco]))["Resumo"]

    assert summary["B2"].value == "Segunda, 28/09/2026"
    assert summary["C2"].value is None


def test_partial_data_is_listed_after_the_notices(api):
    bloco = minimal_block(avisos=["Aviso 1", "Aviso 2"], parcial=True)

    assert open_xlsx(export(api, [bloco]))["Resumo"]["E2"].value == (
        "Aviso 1 | Aviso 2 | Dados parciais"
    )


def test_several_blocks(api):
    blocos = [
        load("ocupacao_profissional"),
        load("pacientes_por_profissional_semana"),
        load("ocupacao_agregada"),
    ]

    response = export(api, blocos)
    workbook = open_xlsx(response)

    assert (
        f'filename="realocai-relatorios-{timezone.localdate():%Y-%m-%d}.xlsx"'
        in (response["Content-Disposition"])
    )
    assert workbook.sheetnames == [
        "Resumo",
        "Por dia (1)",
        "Resumo da semana (1)",
        "Por profissional e dia (2)",
        "Resumo por profissional (2)",
        "Clínica por dia (2)",
        "Por especialidade (3)",
        "Por sala (3)",
    ]
    assert workbook["Resumo"].max_row == 4
    assert workbook["Resumo"]["C3"].value is None  # meta null for patients per professional


def test_sheet_names_are_sanitized_cut_and_unique(api):
    bloco = minimal_block()
    base = copy.deepcopy(bloco["tabelas"][0])
    names = [
        "Ocupação: sala/posto [manhã]?*\\",
        "Um nome de tabela muito comprido que passa de 31",
        "Um nome de tabela muito comprido que passa de 31",
        "resumo",
        "",
    ]
    bloco["tabelas"] = [dict(base, nome=n) for n in names]

    sheetnames = open_xlsx(export(api, [bloco])).sheetnames

    assert sheetnames[0] == "Resumo"
    assert sheetnames[1] == "Ocupação  sala posto  manhã"
    assert sheetnames[2] == "Um nome de tabela muito comprid"
    assert sheetnames[3] == "Um nome de tabela muito compr 2"
    assert sheetnames[4] == "resumo 2"
    assert sheetnames[5] == "Tabela"
    assert all(len(n) <= 31 for n in sheetnames)
    assert len({n.lower() for n in sheetnames}) == len(sheetnames)


def test_suffix_survives_long_names_with_several_blocks(api):
    bloco = minimal_block()
    bloco["tabelas"][0]["nome"] = "x" * 40

    sheetnames = open_xlsx(export(api, [bloco, bloco])).sheetnames

    assert sheetnames[1:] == ["x" * 27 + " (1)", "x" * 27 + " (2)"]


# --- Formula injection --------------------------------------------------------------------


@pytest.mark.parametrize(
    "payload", ["=cmd|' /C calc'!A0", "+1+1", "-1-1", "@SUM(A1)", "\t=1", "\r=1"]
)
def test_formula_like_text_is_stored_as_text(api, payload):
    bloco = minimal_block(titulo=payload, avisos=[payload])
    tabela = bloco["tabelas"][0]
    tabela["colunas"][1]["rotulo"] = payload
    tabela["linhas"][0]["nome"] = payload
    tabela["linhas"][0]["pct"] = payload  # text in a numeric column stays text

    workbook = open_xlsx(export(api, [bloco]))
    summary, sheet = workbook["Resumo"], workbook["Por dia"]

    for cell in (summary["A2"], summary["E2"], sheet["B1"], sheet["B2"], sheet["C2"]):
        assert cell.data_type == "s"
        # XML normalizes a carriage return to a line feed; it is still plain text.
        assert cell.value == payload.replace("\r", "\n")


def test_written_xml_has_no_formula(api):
    import zipfile

    bloco = minimal_block(titulo='=HYPERLINK("http://exemplo.test")')

    response = export(api, [bloco])
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        xml = "".join(
            archive.read(n).decode() for n in archive.namelist() if n.startswith("xl/worksheets/")
        )

    assert "<f>" not in xml


# --- Isolation ----------------------------------------------------------------------------


def test_export_only_uses_the_request_body(api, user):
    other = get_user_model().objects.create_user(
        email="outra.profissional@exemplo.com", password="senha-forte-123"
    )
    secret = minimal_block(titulo="Relatório de outra pessoa")
    conversation = Conversation.objects.create(user=other)
    record_exchange(conversation, "Pergunta", "Resposta", blocos=[secret])

    workbook = open_xlsx(export(api, [minimal_block()]))

    values = [c.value for s in workbook.worksheets for row in s.iter_rows() for c in row]
    assert "Relatório de outra pessoa" not in values
